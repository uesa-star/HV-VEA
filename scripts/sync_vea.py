#!/usr/bin/env python3
"""Descarga el Excel VEA privado de Drive y reemplaza sus tablas en Supabase.

Las credenciales se leen exclusivamente desde variables de entorno de GitHub Actions.
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import unicodedata
from datetime import date, datetime
from typing import Any

import requests
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

SHEETS = {
    "EDAS": "edas",
    "IRAS": "iras",
    "FEBRILES": "febriles",
    "INDIVIDUAL": "individual",
    "SOAT": "soat",
    "VIH": "vih",
    "TBC": "tbc",
}
BATCH = 500


def required(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"Falta la variable/secreto {name}")
    return value


def column_name(value: Any, index: int) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip()).strip("_").lower()
    return text or f"columna_{index + 1}"


def json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value != value:
        return None
    return value


def read_sheet(ws: Any) -> list[dict[str, Any]]:
    rows = ws.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration:
        return []
    names: list[str] = []
    used: dict[str, int] = {}
    for i, header in enumerate(headers):
        name = column_name(header, i)
        used[name] = used.get(name, 0) + 1
        names.append(name if used[name] == 1 else f"{name}_{used[name]}")
    output = []
    for row in rows:
        if not any(value not in (None, "") for value in row):
            continue
        output.append({name: json_value(row[i] if i < len(row) else None) for i, name in enumerate(names)})
    return output


def drive_download(file_id: str, credentials_json: str) -> bytes:
    info = json.loads(credentials_json)
    credentials = service_account.Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/drive.readonly"]
    )
    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    request = service.files().get(fileId=file_id, alt="media")
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buffer.getvalue()


def supabase_request(base: str, key: str, method: str, table: str, **kwargs: Any) -> requests.Response:
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }
    return requests.request(method, f"{base}/rest/v1/{table}", headers=headers, timeout=120, **kwargs)


def replace_table(base: str, key: str, table: str, rows: list[dict[str, Any]]) -> None:
    # Todas las tablas VEA gestionadas por el módulo tienen _row_id; se borra
    # solamente lo que ya existe y luego se carga el nuevo corte completo.
    existing = supabase_request(base, key, "GET", table, params={"select": "_row_id", "limit": "1000000"})
    if existing.status_code >= 300:
        raise RuntimeError(f"{table}: no se pudieron leer filas ({existing.status_code}): {existing.text[:500]}")
    ids = [item.get("_row_id") for item in existing.json() if item.get("_row_id") is not None]
    for start in range(0, len(ids), BATCH):
        part = ids[start:start + BATCH]
        query = ",".join(str(item) for item in part)
        response = supabase_request(base, key, "DELETE", table, params={"_row_id": f"in.({query})"})
        if response.status_code >= 300:
            raise RuntimeError(f"{table}: no se pudieron eliminar filas ({response.status_code}): {response.text[:500]}")
    for start in range(0, len(rows), BATCH):
        payload = rows[start:start + BATCH]
        response = supabase_request(
            base,
            key,
            "POST",
            table,
            headers={
                "apikey": key,
                "Authorization": f"Bearer {key}",
                "Content-Type": "application/json",
                "Prefer": "return=minimal",
            },
            data=json.dumps(payload, ensure_ascii=False),
        )
        if response.status_code >= 300:
            raise RuntimeError(f"{table}: error insertando lote ({response.status_code}): {response.text[:1000]}")
    print(f"{table}: {len(rows)} filas sincronizadas")


def main() -> int:
    file_id = required("GOOGLE_DRIVE_FILE_ID")
    base = required("SUPABASE_URL").rstrip("/")
    key = required("SUPABASE_SERVICE_ROLE_KEY")
    raw = drive_download(file_id, required("GOOGLE_SERVICE_ACCOUNT_JSON"))
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as temp:
        temp.write(raw)
        temp.flush()
        workbook = load_workbook(temp.name, read_only=True, data_only=True)
        parsed: dict[str, list[dict[str, Any]]] = {}
        for sheet, table in SHEETS.items():
            if sheet not in workbook.sheetnames:
                raise RuntimeError(f"No se encontró la hoja obligatoria {sheet}")
            parsed[table] = read_sheet(workbook[sheet])
        workbook.close()
    print("Excel descargado y validado:", ", ".join(f"{k}={len(v)}" for k, v in parsed.items()))
    for table, rows in parsed.items():
        replace_table(base, key, table, rows)
    print("Sincronización VEA completada correctamente.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
